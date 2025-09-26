import pandas as pd
import numpy as np
import rasterio
from rasterio.windows import Window
from sklearn.model_selection import train_test_split
from sklearn.preprocessing import LabelEncoder, StandardScaler
from sklearn.ensemble import RandomForestClassifier
from sklearn.svm import SVC
from sklearn.neighbors import KNeighborsClassifier
from sklearn.linear_model import LogisticRegression
from sklearn.metrics import accuracy_score, classification_report, confusion_matrix
from docx import Document
# Training Data
train_csv = r"D:\Agri_Geospatial\ndvi_3.csv"
df = pd.read_csv(train_csv)
# Features and labels
X = df[["RASTERVALU"]].values
y = df["Type"].values
classes = ["Waterbody", "Settlement", "Barren Land",
           "Moderate Vegetation", "Stressed Vegetation", "Healthy Vegetation"]
le = LabelEncoder()
le.fit(classes)
y_encoded = le.transform(y)
print("Class mapping:", dict(zip(le.classes_, le.transform(le.classes_))))
# Train/test split
X_train, X_test, y_train, y_test = train_test_split(
    X, y_encoded, test_size=0.2, random_state=42, stratify=y_encoded
)
# Define Models 
models = {
    "RF": RandomForestClassifier(
        n_estimators=200, class_weight="balanced", random_state=42, n_jobs=-1
    ),
    "SVM": SVC(
        kernel="rbf", C=10, gamma="scale", class_weight="balanced", random_state=42
    ),
    "KNN": KNeighborsClassifier(n_neighbors=3, weights="distance", n_jobs=-1),
    "LR": LogisticRegression(
        multi_class="multinomial", solver="lbfgs", max_iter=500, class_weight="balanced"
    ),
}
scaler = StandardScaler()
scaler.fit(X_train)
# Rasters
input_raster = r"D:\Agri_Geospatial\Prisma_resampled\Study_area3\Prisma_downscaled_area3_NDVI.tif"
output_dir = r"D:\Agri_Geospatial"
# Train, Evaluate, Report, Apply Raster 
for name, model in models.items():
    print(f"\nTraining {name}")
    # Scale features for SVM, KNN, LR
    if name in ["SVM", "KNN", "LR"]:
        X_train_scaled = scaler.transform(X_train)
        X_test_scaled = scaler.transform(X_test)
        model.fit(X_train_scaled, y_train)
        y_pred = model.predict(X_test_scaled)
    else:
        model.fit(X_train, y_train)
        y_pred = model.predict(X_test)

    # Evaluation
    acc = accuracy_score(y_test, y_pred)
    cm = confusion_matrix(y_test, y_pred, labels=range(len(classes)))
    cr = classification_report(y_test, y_pred, target_names=le.classes_)
    print(f"\n{name} Accuracy: {acc:.4f}")
    print("\nConfusion Matrix:\n", cm)
    print("\nClassification Report:\n", cr)
    # Report 
    report_doc = fr"{output_dir}\Prisma_{name}_LULC_NDVI_area3_Report.docx"
    doc = Document()
    doc.add_heading(f"Prisma {name} LULC Classification (NDVI_area3)", 0)
    doc.add_paragraph(f"Accuracy: {acc:.4f}")
    doc.add_paragraph("Class Mapping:\n" + str(dict(zip(le.classes_, le.transform(le.classes_)))))
    doc.add_paragraph("Confusion Matrix:\n" + str(cm))
    doc.add_paragraph("Classification Report:\n" + cr)
    doc.save(report_doc)
    print(f"Word report saved → {report_doc}")

    # Apply Model to Raster
    output_raster = fr"{output_dir}\Prisma_classified_{name}_NDVI_area3.tif"
    with rasterio.open(input_raster) as src:
        meta = src.meta.copy()
        meta.update({"count": 1, "dtype": "uint8", "compress": "lzw", "nodata": 0})
        with rasterio.open(output_raster, "w", **meta) as dst:
            block_size = 1024
            for i in range(0, src.height, block_size):
                for j in range(0, src.width, block_size):
                    win = Window(j, i, min(block_size, src.width - j), min(block_size, src.height - i))
                    block = src.read(1, window=win).astype(np.float32)
                    # Flatten and replace NaN with 0
                    flat = block.reshape(-1, 1)
                    flat = np.nan_to_num(flat, nan=0.0)
                    # Scale if needed
                    if name in ["SVM", "KNN", "LR"]:
                        flat = scaler.transform(flat)
                    # Predict
                    pred = model.predict(flat) + 1
                    classified_block = pred.reshape(block.shape)
                    # NoData for NaNs
                    classified_block[np.isnan(block)] = 0
                    dst.write(classified_block.astype(np.uint8), 1, window=win)
    print(f"{name} classified map saved → {output_raster}")
